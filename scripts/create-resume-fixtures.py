from pathlib import Path
import argparse

from docx import Document
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "tests" / "fixtures"
OUT.mkdir(parents=True, exist_ok=True)

FONT_PATH = Path(r"C:\Windows\Fonts\msjh.ttc")
pdfmetrics.registerFont(TTFont("FixtureSans", str(FONT_PATH), subfontIndex=0))


def draw_lines(pdf, lines, start_y=800):
    pdf.setFont("FixtureSans", 11)
    y = start_y
    for line in lines:
        if line == "---PAGE---":
            pdf.showPage()
            pdf.setFont("FixtureSans", 11)
            y = 800
            continue
        pdf.drawString(58, y, line)
        y -= 22


single_lines = [
    "Avery Lin",
    "Frontend Developer",
    "Summary",
    "Builds accessible web products and collaborates with product teams.",
    "Technical Skills",
    "HTML5, CSS3, JS, TS, React.js, Tailwind CSS, Git, GitHub, RESTful APIs",
    "Work Experience",
    "Bright Harbor Studio - Frontend Developer",
    "2024/01 - Present",
    "Built React dashboards and integrated REST APIs.",
    "Projects",
    "Job Quest Guild - React, TypeScript, Supabase, PostgreSQL",
    "Education",
    "National Example University - Computer Science",
]

multi_lines = [
    "Jordan Chen",
    "Full Stack Developer",
    "Skills",
    "Vue, NodeJS, Express, Postgres, Docker, AWS, GitLab",
    "Experience",
    "Northwind Systems - Web Engineer",
    "2022 - 2024",
    "Maintained customer portals and backend services.",
    "---PAGE---",
    "Projects",
    "Life Tracker - Vue, Node.js, PostgreSQL",
    "Inventory Portal - Java, Spring Boot, MySQL",
    "Education",
    "Example College - Information Technology",
]

def create_pdfs():
    pdf = canvas.Canvas(str(OUT / "resume-text.pdf"), pagesize=A4)
    draw_lines(pdf, single_lines)
    pdf.save()

    pdf = canvas.Canvas(str(OUT / "resume-multipage.pdf"), pagesize=A4)
    draw_lines(pdf, multi_lines)
    pdf.save()

    image_path = OUT / "scanned-resume-source.png"
    image = Image.new("RGB", (1240, 1754), "white")
    draw = ImageDraw.Draw(image)
    draw.text((110, 140), "SCANNED RESUME IMAGE - NO EMBEDDED TEXT", fill="black")
    draw.text((110, 230), "This fixture verifies the no-OCR warning.", fill="black")
    image.save(image_path)
    pdf = canvas.Canvas(str(OUT / "resume-scanned.pdf"), pagesize=A4)
    pdf.drawImage(str(image_path), 0, 0, width=A4[0], height=A4[1])
    pdf.save()
    image_path.unlink()


def create_docx():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    title = doc.add_paragraph(style="Title")
    title.add_run("林怡君")
    doc.add_paragraph("前端工程師")
    for heading, body in [
        ("自我介紹", "專注於可維護的前端產品，具備跨團隊協作與 API 串接經驗。"),
        ("技能", "HTML、CSS、JavaScript、TS、ReactJS、Vite、Git、RESTful API、Supabase、Postgres"),
        ("工作經驗", "星河數位 前端工程師\n2023/06 - 至今\n開發 React 後台並串接 REST API。"),
        ("專案", "Portfolio Guild - React.js、TypeScript、Supabase\nData Console - Java、Spring、PostgreSQL"),
        ("學歷", "範例科技大學 資訊工程系"),
    ]:
        doc.add_heading(heading, level=1)
        for line in body.split("\n"):
            doc.add_paragraph(line)

    for style_name in ["Normal", "Title", "Heading 1"]:
        style = doc.styles[style_name]
        style.font.name = "Microsoft JhengHei"
        style.font.size = Pt(11 if style_name == "Normal" else 16)
    doc.save(OUT / "resume-text.docx")


parser = argparse.ArgumentParser()
parser.add_argument("mode", choices=["pdf", "docx"])
args = parser.parse_args()
create_pdfs() if args.mode == "pdf" else create_docx()
print(OUT)
